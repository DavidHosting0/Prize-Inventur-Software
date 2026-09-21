# -*- coding: utf-8 -*-
"""Generate non-technical sales Word document for Prize Hotel."""
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

OUT = r"c:\Users\ytmad\Desktop\Prize Inventur Software\Prize-Hotel-Produktvorstellung.docx"


def main() -> None:
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0x2C, 0x2C, 0x2C)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    for i in range(1, 4):
        s = doc.styles[f"Heading {i}"]
        s.font.name = "Calibri"
        s.font.color.rgb = RGBColor(0x1A, 0x3A, 0x5C)
        s.font.bold = True
        if i == 1:
            s.font.size = Pt(20)
            s.paragraph_format.space_before = Pt(18)
            s.paragraph_format.space_after = Pt(10)
        elif i == 2:
            s.font.size = Pt(14)
            s.paragraph_format.space_before = Pt(14)
            s.paragraph_format.space_after = Pt(6)
        else:
            s.font.size = Pt(12)
            s.paragraph_format.space_before = Pt(10)
            s.paragraph_format.space_after = Pt(4)

    def bullets(items: list[str]) -> None:
        for item in items:
            p = doc.add_paragraph(item, style="List Bullet")
            p.paragraph_format.space_after = Pt(3)

    # --- Title page ---
    for _ in range(3):
        doc.add_paragraph()

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("PRIZE HOTEL")
    r.bold = True
    r.font.size = Pt(32)
    r.font.color.rgb = RGBColor(0x1A, 0x3A, 0x5C)
    r.font.name = "Calibri"

    st = doc.add_paragraph()
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = st.add_run("Inventur & Kasse")
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor(0x3D, 0x6B, 0x8F)
    r.font.name = "Calibri"

    doc.add_paragraph()

    tag = doc.add_paragraph()
    tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = tag.add_run(
        "Die smarte Lösung für Lager, Inventur, Einkauf und Verkauf\n"
        "im Hotel – klar, schnell und aus einer Hand."
    )
    r.font.size = Pt(13)
    r.font.italic = True
    r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    doc.add_paragraph()

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Produktvorstellung für Entscheider")
    r.font.size = Pt(12)
    r.bold = True

    datum = doc.add_paragraph()
    datum.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = datum.add_run(date.today().strftime("%d.%m.%Y"))
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    doc.add_page_break()

    # --- 1 Problem ---
    doc.add_heading("1. Das Problem – was Hotels täglich kostet", level=1)
    doc.add_paragraph(
        "Viele Hotels steuern Lager und Verkauf noch mit Excel-Listen, Zettelwirtschaft und "
        "getrennten Systemen. Das führt zu Unsicherheit, Zeitverlust und unnötigen Kosten – "
        "besonders in Bar, Restaurant und Frühstück."
    )
    bullets(
        [
            "Inventuren dauern lange, sind fehleranfällig und oft veraltet, bevor sie fertig sind.",
            "Niemand weiß genau, was wirklich im Lager steht – und was schon längst fehlt.",
            "Food Cost und Schwund bleiben undurchsichtig, bis der Monatsabschluss kommt.",
            "Nachbestellungen passieren zu spät oder zu früh – entweder Out-of-Stock oder zu viel Kapital im Regal.",
            "Kasse und Lager sprechen nicht miteinander: Verkäufe werden gebucht, Bestand nicht oder nur manuell.",
            "Bei mehreren Hotels fehlt der Konzernblick: jedes Haus rechnet anders.",
        ]
    )
    doc.add_paragraph(
        "Das Ergebnis: höhere Kosten, gestresste Teams und Entscheidungen ohne belastbare Zahlen."
    )

    # --- 2 Solution ---
    doc.add_heading("2. Die Lösung – Prize Hotel", level=1)
    doc.add_paragraph(
        "Prize Hotel ist die zentrale Betriebsplattform für Inventur, Lager, Einkauf, F&B und Kasse. "
        "Alles, was Ihr Haus für den Wareneinsatz braucht, läuft in einem System – verständlich für "
        "das Team vor Ort und transparent für die Direktion."
    )
    doc.add_paragraph(
        "Jedes Hotel arbeitet mit einem klaren, zentralen Lager. Verkäufe, Wareneingänge, Inventuren "
        "und Verluste fließen dort zusammen. So haben Sie jederzeit einen echten Überblick – "
        "nicht erst am Monatsende."
    )
    p = doc.add_paragraph()
    r = p.add_run("Ein System. Ein Lager. Klare Zahlen.")
    r.bold = True

    # --- 3 Benefits ---
    doc.add_heading("3. Nutzen & Mehrwert", level=1)
    bullets(
        [
            "Weniger Schwund und weniger Überraschungen beim Food Cost",
            "Schnellere, zuverlässigere Inventuren – auch mit Scanner",
            "Bestand, der mit der Kasse mitgeht: verkauft = abgebucht",
            "Bessere Nachbestellungen dank Bestellvorschlägen und Lieferantenstammdaten",
            "Klare Rechte: jeder sieht und darf, was zu seiner Rolle passt",
            "Berichte für Controlling – exportierbar für Excel, Präsentationen und Abschluss",
            "Konzernsteuerung für Hotelgruppen: alle Häuser im Blick",
        ]
    )
    doc.add_paragraph(
        "Kurz gesagt: Prize Hotel spart Zeit im Alltag, schützt Ihre Margen und gibt Ihnen die "
        "Kontrolle zurück – ohne dass Ihr Team eine Schulung zum IT-Spezialisten braucht."
    )

    # --- 4 Features ---
    doc.add_heading("4. Was die Software kann – Modul für Modul", level=1)
    doc.add_paragraph(
        "Nachfolgend die wichtigsten Bereiche in der Sprache des Hotelbetriebs – ohne Technik-Fachchinesisch."
    )

    doc.add_heading("4.1 Dashboard – der Tagesstart", level=2)
    doc.add_paragraph(
        "Das Dashboard zeigt auf einen Blick, wie das Haus steht: Kennzahlen, Trends und den "
        "schnellen Einstieg in die Kasse. Ideal für Schichtbeginn und kurze Lagebesprechungen."
    )

    doc.add_heading("4.2 Bestand / Lager – Ihre zentrale Wahrheit", level=2)
    doc.add_paragraph(
        "Im Lager sehen Sie Mengen, Werte und Status (z. B. kritisch, niedrig, in Ordnung). "
        "Sie können Bestände manuell anpassen und jede Bewegung nachverfolgen – Verkäufe, "
        "Lieferungen, Korrekturen und Verluste."
    )
    doc.add_paragraph(
        "Ein Lager pro Hotel bedeutet: keine Verwirrung zwischen mehreren Bestandsorten, "
        "klare Verantwortung und einfache Inventur."
    )

    doc.add_heading("4.3 Inventur – zählen, prüfen, abschließen", level=2)
    doc.add_paragraph(
        "Starten Sie eine Inventur, zählen Sie Artikel manuell oder per Scanner, sehen Sie "
        "Differenzen sofort und schließen Sie die Inventur kontrolliert ab. Optional mit "
        "Prüfschritt, bevor Korrekturen wirksam werden."
    )
    bullets(
        [
            "Übersicht: ungezählt, alle Positionen, nur Differenzen",
            "Geeignet auch für Flüssigkeiten (z. B. volle und angebrochene Flaschen)",
            "Regeln pro Hotel einstellbar – so, wie Ihr Haus arbeitet",
        ]
    )
    doc.add_paragraph(
        "So wird Inventur vom gefürchteten Pflichttermin zum sauberen, wiederholbaren Prozess."
    )

    doc.add_heading("4.4 Wareneingang – vom Lieferschein ins Lager", level=2)
    doc.add_paragraph(
        "Erfassen Sie Lieferungen manuell oder fotografieren Sie den Lieferschein: Die Software "
        "erkennt die Positionen intelligent aus den Fotos, Ihr Team prüft und bucht mit einem "
        "Klick ins Lager."
    )
    bullets(
        [
            "Weniger Tipparbeit, weniger Erfassungsfehler",
            "Entwurf speichern und erst nach Kontrolle bestätigen",
            "Nachvollziehbare Historie jeder Lieferung",
        ]
    )

    doc.add_heading("4.5 Bestellungen & Lieferanten", level=2)
    doc.add_paragraph(
        "Aus Mindest- und Optimalbeständen entstehen Bestellvorschläge. Daraus legen Sie "
        "Bestellungen an und steuern Ihre Lieferanten mit Kontaktdaten, Lieferzeiten und "
        "Mindestbestellwerten – alles an einem Ort."
    )

    doc.add_heading("4.6 Artikel & Rezepte", level=2)
    doc.add_paragraph(
        "Im Artikelstamm pflegen Sie Produkte inkl. Preisen, Mehrwertsteuer, Barcodes und "
        "Lagergrenzen. Rezepte verbinden Zutaten zu verkaufbaren Gerichten und Getränken – "
        "inkl. Food-Cost-Sicht und theoretischem Verbrauch."
    )
    doc.add_paragraph(
        "Das bedeutet: Wenn an der Kasse ein Cocktail verkauft wird, wissen Sie nicht nur den "
        "Umsatz – sondern auch, welche Zutaten das Lager verlassen haben."
    )

    doc.add_heading("4.7 Kasse / Verkauf – schnell, klar, bestandswirksam", level=2)
    doc.add_paragraph(
        "Das Verkaufsterminal ist für den Betrieb gemacht: Favoriten, Kategorien, Barcode, "
        "Rabatte, Gutscheine, Gratisartikel, Parken, Storno und Rückerstattung. Zahlung per "
        "Bargeld, Karte, TWINT und weitere Wege."
    )
    bullets(
        [
            "Bestand wird erst bei Bezahlung reduziert – nicht schon beim Parken",
            "Kassenabschluss für die Schicht",
            "Belege drucken oder als PDF speichern",
            "Kein Zimmerfolio nötig: ideal für Bar, Outlet und Nebenverkäufe",
        ]
    )
    doc.add_paragraph(
        "Im Kassenkonfigurator richten Sie Verkaufsartikel, Kategorien, Preise und Layout ein – "
        "getrennt vom reinen Lagerstamm, damit die Kasse übersichtlich bleibt."
    )

    doc.add_heading("4.8 Food Waste, Frühstück & Minibar", level=2)
    doc.add_paragraph(
        "Food Waste: Verluste mit Gründen erfassen (z. B. abgelaufen, Überproduktion, "
        "Buffet-Rest) – für Transparenz und Gegensteuerung."
    )
    doc.add_paragraph(
        "Frühstück: Gästezahlen, Kosten und Prognose unterstützen Planung und Controlling des Buffets."
    )
    doc.add_paragraph(
        "Minibar: Verbrauch pro Zimmer erfassen und Bestand nachführen – ohne automatische "
        "Zimmerbelastung. Ideal, wenn die Abrechnung ohnehin über ein anderes System läuft "
        "oder bewusst manuell bleibt."
    )

    doc.add_heading("4.9 Berichte & Auswertungen", level=2)
    doc.add_paragraph(
        "Zusammenfassungen zu Umsatz, Verlusten, Bewegungen und kritischen Beständen – "
        "exportierbar als Tabelle oder PDF für Meetings, Revision und Monatsabschluss."
    )

    doc.add_heading("4.10 Benutzer & Rechte", level=2)
    doc.add_paragraph(
        "Jeder Mitarbeitende sieht nur das, was zur Rolle passt: Bar bedient die Kasse, Lager "
        "steuert Bestand und Inventur, F&B und Direktion behalten den Überblick. Rechte sind "
        "klar und nachvollziehbar."
    )

    doc.add_heading("4.11 Konzern- & Gruppensteuerung", level=2)
    doc.add_paragraph(
        "Für Hotelgruppen: Portfolio-Dashboard, Hotelverwaltung, Analysen über mehrere Häuser, "
        "zentrale Benutzer und Berechtigungen sowie ein Audit-Protokoll. Neue Hotels lassen "
        "sich strukturiert anlegen und sofort betriebsbereit vorbereiten."
    )
    doc.add_paragraph(
        "Optional unterstützen zentrale Einstellungen intelligente Funktionen wie die "
        "Lieferschein-Erkennung – einmal konfiguriert, für die ganze Gruppe nutzbar."
    )

    # --- 5 Who ---
    doc.add_heading("5. Für wen ist Prize Hotel gemacht?", level=1)

    doc.add_heading("Hoteldirektion & General Management", level=3)
    doc.add_paragraph(
        "Kontrolle über Kosten, Transparenz und Berichte – ohne im Detail mitzählen zu müssen."
    )

    doc.add_heading("F&B-Management", level=3)
    doc.add_paragraph(
        "Food Cost, Waste, Frühstück und Rezepte im Griff – Entscheidungen auf Basis echter Zahlen."
    )

    doc.add_heading("Bar, Restaurant & Outlet", level=3)
    doc.add_paragraph(
        "Schnelle Kasse, klare Favoriten, weniger Papier – und Bestand, der mitläuft."
    )

    doc.add_heading("Lager & Einkauf", level=3)
    doc.add_paragraph(
        "Inventur, Wareneingang, Bestellungen und Lieferanten in einem durchgängigen Ablauf."
    )

    doc.add_heading("Konzern / Hotelgruppe", level=3)
    doc.add_paragraph(
        "Vergleichbarkeit zwischen Häusern, zentrale Steuerung und einheitliche Prozesse."
    )

    # --- 6 Scenarios ---
    doc.add_heading("6. So sieht der Alltag aus", level=1)

    doc.add_heading("Szenario A – Inventurtag", level=2)
    doc.add_paragraph(
        "Das Team startet eine Inventur, zählt mit Scanner oder am Bildschirm, sieht Differenzen "
        "sofort und schließt nach kurzer Prüfung ab. Der Bestand ist aktualisiert – ohne "
        "wochenlange Excel-Nacharbeit."
    )

    doc.add_heading("Szenario B – Schicht an der Bar", level=2)
    doc.add_paragraph(
        "Die Bar öffnet die Kasse, verkauft Favoriten und Rezeptgetränke, nimmt Zahlungen "
        "entgegen und macht am Ende den Kassenabschluss. Jeder bezahlte Verkauf hat den "
        "Bestand korrekt reduziert."
    )

    doc.add_heading("Szenario C – Lieferung kommt", level=2)
    doc.add_paragraph(
        "Der Lieferschein wird fotografiert oder die Positionen manuell erfasst. Nach kurzer "
        "Kontrolle wird gebucht – das Lager ist aktuell, bevor die Ware im Regal steht."
    )

    doc.add_heading("Szenario D – Blick der Gruppe", level=2)
    doc.add_paragraph(
        "Die Konzernleitung vergleicht Häuser, erkennt Ausreißer bei Waste oder Bestand und "
        "steuert nach – mit einer gemeinsamen Sprache und denselben Prozessen."
    )

    # --- 7 Usage ---
    doc.add_heading("7. Sprachen & Nutzung", level=1)
    bullets(
        [
            "Oberfläche auf Deutsch und Englisch – ideal für internationale Teams",
            "Nutzung im Browser am PC oder Laptop im Büro und an der Station",
            "Auch als App auf dem Handy nutzbar – für unterwegs und im Lager",
            "Optional als Desktop-Anwendung – dieselbe vertraute Oberfläche",
        ]
    )
    doc.add_paragraph(
        "Ihr Team arbeitet dort, wo der Betrieb stattfindet – nicht nur am Büro-PC."
    )

    # --- 8 CTA ---
    doc.add_heading("8. Nächste Schritte", level=1)
    doc.add_paragraph(
        "Prize Hotel ist bereit, Ihr Haus oder Ihre Gruppe zu unterstützen. Der einfachste Weg:"
    )
    bullets(
        [
            "Kurzdemo mit Ihren typischen Abläufen (Inventur, Kasse, Wareneingang)",
            "Pilot in einem Hotel – messbarer Nutzen in wenigen Wochen",
            "Rollout auf weitere Häuser mit einheitlichen Standards",
        ]
    )

    doc.add_paragraph()
    cta = doc.add_paragraph()
    cta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cta.add_run(
        "Lassen Sie uns gemeinsam zeigen, wie viel Zeit und Kosten Sie sparen können."
    )
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0x1A, 0x3A, 0x5C)

    doc.add_paragraph()
    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("Prize Hotel – Inventur & Kasse\nProduktvorstellung")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
